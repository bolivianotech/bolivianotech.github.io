#!/usr/bin/env python
"""Genera los CVs de Jimmy Requena (DOCX + PDF) en formato ATS.

Uso:
    python cv/build_cv.py                      # todos los perfiles, ES y EN
    python cv/build_cv.py general              # un perfil, ES y EN
    python cv/build_cv.py docente es           # un perfil, un idioma
    python cv/build_cv.py general --no-pdf

Fuente unica de datos: cv/data/master.json. Los perfiles viven en cv/profiles/profiles.json.
Los perfiles SOLO eligen y ordenan contenido del maestro; nunca agregan hechos nuevos.
"""
import json
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MASTER = json.loads((ROOT / "data" / "master.json").read_text(encoding="utf-8-sig"))
PROFILES = json.loads((ROOT / "profiles" / "profiles.json").read_text(encoding="utf-8-sig"))
OUT = ROOT / "output"
FONT = "Calibri"
SOFFICE = [
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "soffice",
]


def tr(value, lang):
    """Resuelve un campo que puede ser str o {es, en}."""
    return value[lang] if isinstance(value, dict) else value


def select_bullets(job, profile):
    limit = profile["max_bullets"].get(job["id"], profile["max_bullets"]["default"])
    prio = profile["priority_tags"]
    scored = []
    for idx, b in enumerate(job["bullets"]):
        score = sum(len(prio) - prio.index(t) for t in b["tags"] if t in prio)
        scored.append((-score, idx, b))
    chosen = sorted(scored)[:limit]
    return [b for _, _, b in sorted(chosen, key=lambda x: x[1])]  # conserva orden original


def set_style(doc):
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(2)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.08
    h = doc.styles["Heading 1"]
    h.font.name = FONT
    h.font.size = Pt(11.5)
    h.font.bold = True
    h.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(3)
    h.paragraph_format.keep_with_next = True
    lb = doc.styles["List Bullet"]
    lb.font.name = FONT
    lb.font.size = Pt(10.5)
    lb.paragraph_format.space_after = Pt(1.5)


def bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F2A44")
    borders.append(bottom)
    pPr.append(borders)


def heading(doc, text):
    p = doc.add_heading(text, level=1)
    bottom_border(p)
    return p


def para(doc, text="", bold=False, italic=False, size=None, align=None, space_after=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    return p


def build(profile_id, lang):
    profile = PROFILES[profile_id]
    lab = MASTER["section_labels"][lang]
    c = MASTER["contact"]
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)  # carta/letter
    s.left_margin = s.right_margin = Inches(0.8)
    s.top_margin = s.bottom_margin = Inches(0.55)
    set_style(doc)
    doc.core_properties.title = f"CV {MASTER['name']}"
    doc.core_properties.author = MASTER["name"]

    # Cabecera (en el cuerpo, no en el header del documento: los ATS leen mejor asi)
    p = para(doc, MASTER["name"].upper(), bold=True, size=18, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    p.runs[0].font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
    para(doc, tr(profile["headline"], lang), bold=True, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    para(doc, f"{c['location']}  |  {c['phones'][0]}  |  {c['phones'][1]}  |  {c['email']}",
         size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    para(doc, f"LinkedIn: {c['linkedin']}  |  {lab['web']}: {c['web']}",
         size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)

    heading(doc, lab["summary"])
    para(doc, tr(profile["summary"], lang))

    heading(doc, lab["skills"])
    groups = {g["id"]: g for g in MASTER["skills"]}
    for gid in profile["skill_order"]:
        g = groups[gid]
        p = doc.add_paragraph()
        p.add_run(f"{tr(g['label'], lang)}: ").bold = True
        p.add_run(", ".join(tr(i, lang) for i in g["items"]))
    # Pruebas y calidad solo cuando el perfil la prioriza
    if "qa" not in profile["skill_order"] and "qa" in profile["priority_tags"]:
        g = groups["qa"]
        p = doc.add_paragraph()
        p.add_run(f"{tr(g['label'], lang)}: ").bold = True
        p.add_run(", ".join(tr(i, lang) for i in g["items"]))

    heading(doc, lab["experience"])
    jobs = {j["id"]: j for j in MASTER["experience"]}
    for jid in profile["experience_order"]:
        j = jobs[jid]
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.keep_with_next = True
        p.add_run(f"{tr(j['role'], lang)}").bold = True
        p.add_run(f" — {j['org']}").bold = True
        bits = [b for b in [j["place"]] if b]
        if j["start"] or j["end"]:
            end = tr(j["end"], lang) if j["end"] else ""
            bits.append(f"{j['start']} – {end}".strip(" –"))
        if bits:
            q = para(doc, "  |  ".join(bits), italic=True, size=9.5, space_after=1)
            q.paragraph_format.keep_with_next = True
        for b in select_bullets(j, profile):
            doc.add_paragraph(tr(b, lang), style="List Bullet")

    # "Otra experiencia": omite lo que el perfil ya detalla arriba
    extra = [o for o in MASTER["other_experience"] if o["id"] not in profile["experience_order"]]
    if extra:
        heading(doc, lab["other"])
        for o in extra:
            doc.add_paragraph(o[lang], style="List Bullet")

    if profile.get("teaching_section"):
        heading(doc, lab["teaching"])
        para(doc, MASTER["teaching_history"][lang])

    heading(doc, lab["education"])
    for e in MASTER["education"]:
        doc.add_paragraph(tr(e, lang), style="List Bullet")

    heading(doc, lab["certs"])
    para(doc, "  ·  ".join(MASTER["certifications"]))

    heading(doc, lab["languages"])
    para(doc, MASTER["languages"][lang])

    heading(doc, lab["refs"])
    para(doc, lab["refs_text"])

    OUT.mkdir(exist_ok=True)
    d = OUT / profile_id
    d.mkdir(exist_ok=True)
    path = d / f"CV_Jimmy_Requena_{profile_id}_{lang.upper()}.docx"
    doc.save(path)
    return path


def to_pdf(docx_path):
    for exe in SOFFICE:
        if exe == "soffice" or Path(exe).exists():
            subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                           check=True, capture_output=True)
            return docx_path.with_suffix(".pdf")
    raise RuntimeError("No se encontro LibreOffice (soffice) para generar el PDF")


def main(argv):
    pdf = "--no-pdf" not in argv
    args = [a for a in argv if not a.startswith("--")]
    pids = [args[0]] if args else list(PROFILES)
    langs = [args[1]] if len(args) > 1 else ["es", "en"]
    for pid in pids:
        if pid not in PROFILES:
            sys.exit(f"Perfil desconocido: {pid}. Disponibles: {', '.join(PROFILES)}")
        for lang in langs:
            f = build(pid, lang)
            out = to_pdf(f) if pdf else f
            print("OK", out)
    # Los CVs 'general' son los que se publican en la pagina web
    if "general" in pids and pdf:
        for lang in langs:
            src = OUT / "general" / f"CV_Jimmy_Requena_general_{lang.upper()}.pdf"
            shutil.copyfile(src, ROOT / f"CV_Jimmy_Requena_{lang.upper()}.pdf")
            print("PUBLICADO", ROOT / f"CV_Jimmy_Requena_{lang.upper()}.pdf")


if __name__ == "__main__":
    main(sys.argv[1:])
